# 엑셀 기획표(webtoon/planning/오늘도우리집_기획표.xlsx)를 만든다.
# 사용법(저장소 루트에서):
#   node webtoon/tools/export-data.cjs data.json
#   python3 webtoon/tools/make-planning-xlsx.py data.json "webtoon/planning/오늘도우리집_기획표.xlsx"
# 필요한 것: python3, openpyxl (pip install openpyxl)
# 새 화를 만들면 아래 POINTS 에 그 화의 ‘웃음 포인트 / 감동 포인트’를 한 줄 더한다.
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

d = json.load(open(sys.argv[1]))
out = sys.argv[2]
NAMES = {c["id"]: c["name"] for c in d["CHARACTERS"]}
BG = {"kitchen": "부엌", "living": "거실", "living-night": "거실(밤)", "halmae": "할머니 댁", "bedroom": "방",
      "bedroom-night": "방(밤)", "entrance": "현관", "street": "퇴근길", "burst": "효과: 집중선", "sparkle": "효과: 반짝",
      "gloom": "효과: 어두움", "speed": "효과: 속도선", "memory": "회상", "plain": "단색"}
LOOK = {  # 아이 그림에서 가져온 특징
    "jaei": ("머리띠 + 연한 동그라미 올림머리 + 옆으로 묶은 머리 (아이 그림 그대로)", "분홍 원피스, 흰 동그란 깃 · 입버릇 “아니야!” · 정리왕, 숙제 미리 끝 · 치즈 극혐 · 태오랑 같이 마블(스파이더맨·판타스틱 4·어벤져스) 좋아함"),
    "halmeoni": ("뽀글뽀글 파마머리, 눈가 주름", "연보라 꽃무늬 블라우스, 자주색 치마"),
    "eomma": ("정수리 똥머리, 분홍 머리끈", "연분홍 원피스, 리본 허리띠"),
    "appa": ("짧고 단정한 머리, 귀", "회색 셔츠, 갈색 벨트 (출근할 땐 넥타이)"),
    "harabeoji": ("훤한 정수리 + 옆머리, 짙은 눈썹, 눈가 주름", "갈색 카디건, 베이지 바지"),
    "taeo": ("짧은 머리, 삐죽 솟은 한 가닥, 귀", "파란 윗옷·바지 (태권도복: 파란띠에 초록 줄) · 트레이드마크: 늘 들고 다니며 핥아 먹는 반쪽 레몬 (신 과일을 제일 좋아함)"),
}
POINTS = {
    1: ("아빠가 “배불러서 괜찮다”면서 손은 딸기로 / 아빠가 자른 네 조각은 크기가 제각각", "엄마는 늘 가장 작은 조각을 가져간다는 걸 태오가 알아챈다"),
    2: ("“안 보이면 없는 거야” 소파 밑에 숨긴 그림 / 마지막에 아빠도 화분을 고백", "“알고 있었어. 네가 먼저 말해 주길 기다렸어.”"),
    3: ("“충전? 아빠가 핸드폰이야?” / 자는 아빠의 방귀로 ‘충전 완료’", "아이들이 만든 쿠폰에 아빠 배터리가 100%로 찬다"),
    4: ("새까맣게 탄 계란 / 빨간 양말 때문에 분홍이 된 아빠 셔츠 / 짠 죽", "짠 죽을 “세상에서 제일 맛있다”고 말하는 엄마"),
    5: ("영상통화 화면에 할아버지 머리만 번쩍 / 할머니는 파마만 보인다", "아기 재이에게 숟가락질을 백 번 가르쳐 준 할아버지"),
    8: ("“아빠 내 말 안 들으면 피노키오야!” 코를 만지는 아빠 / “맛있는 거 안 주면 소리 지른다!” → 쩌렁쩌렁", "“아빠가 먼저 약속을 안 지켰네. 미안해.” / 마법의 말 ‘주세요’"),
    9: ("머리끈 뺏기·흉내·낙서·방귀로 줄어드는 누나 참을성 배터리 / “김태오!!!” 엉덩이 찰싹", "혼난 태오를 제일 먼저 안아 주는 누나 / “누나가 좋아서 그랬어.”"),
    10: ("“패드로 어벤져스 게임 하자!” (보드게임) / “늘어나는 애 이름 뭐야?” 또 묻기 / “오구마꼬치 맞아!! 선샘민도 그랬어!!” 우기다 긁적긁적", "속삭인 말도 다 듣는 태오 / “글씨는 몰라도 가족 마음 읽는 건 1등이야.”"),
    11: ("“블루베리 출동! 쉬터맨도 출동!” / 태권도 발차기하다 삐끗", "다리 부러진 초록 영웅에게 반창고를 붙여 주는 태오 / “다리 아파도 영웅이야!”"),
    12: ("“어벤져스는 가년이야!!” 우기기 / 아빠 “스티브 로저스” → “땡! 스테잌로저야!” / 할아버지 얼떨결에 “우엉?” 정답", "할머니 “어벤져스에 가면… 우리 태오도 있고!” / 우리 집 어벤져스는 여섯 명"),
    13: ("태오가 제일 좋아하는 레몬·자두·파인애플 / “아빠 레몬 먹어!!” “싫어!” “싫어하는 것도 먹어야 튼튼해지고 근육도 생겨!” — 아빠 말 그대로 돌려주기 / 시큼!!! 부르르 → 태오 “뭐가 셔? 맛있는데.”", "아빠가 먼저 레몬을 먹자 태오도 브로콜리 도전 / 그리고 “누나 레몬 먹어!” → “김태오!!”"),
    14: ("“아빠 레스테스통 알아?” “레스…테스…통? 아빠 머리통?” / 누나 통역: 레시피 스톤! / “그거 아빠가 한 말인데?”", "할머니의 레시피 스톤 = 오래된 요리 공책 / “그건 레스테스통이 아니라 할미 사랑이여~”"),
    15: ("아빠가 티비 켜자마자 “내꺼 틀어 줘!!” / 밥 먹고 10분 뒤 “심심해~ 맛있는 거 줘!” / “싸우자!” 건틀렛 낀 타노스 아빠 “아임 인에비터블!” → 태오 “아임 인에터블!!”", "열 번 진 아빠를 보고 태오가 타노스를 맡아 일부러 져 준다 / “아빠 맨날 졌잖아. 오늘은 아빠가 이겨.”"),
    16: ("“너라고 하지 마라! 나도 엄마 아빠한테 너라고 한다!” / “재이 아니야!! 재이 누나야!!” / 엄마가 할머니를 “엄마!” → “엄마 아니야!! 할머니야!!” / “김서방 아니야! 아빠야!”", "할머니가 그려 준 가족 나무 / “나 이름 엄청 많다!” — 아들, 동생, 손자 / 그리고 “태오 님이야!” → “김태오!!”"),
    17: ("숙제 미리 끝, 방 정리 끝 → 태오가 와르르 “아니야! 아니야! 아니야!” / “맛있는데… 지금은 먹기 싫어.” / 치즈 듬뿍 피자 “치즈 냄새!! 극혐!!” / “아니야인데 고마워?”", "태오가 “누나 쪽은 치즈 빼 줘!” — 누나가 싫어하는 걸 다 기억한 동생 / 티니핑도 위시캣도 둘 다"),
    18: ("게임에 진 태오의 “바보야! 피노키오야! 때릴 거야! 침 뱉을 거야!”가 화살이 되어 날아간다 / 베개 펀치 “쾅쾅” / 아빠 “나도 피노키오 아니지?” “아빠는… 타노스야.”", "혼내는 대신 “화났구나, 화나는 건 괜찮아.” / 우리 집 주문 “나 화났어!” 후~ / “누나는 바보 아니야, 블랙 위도우야.”"),
    19: ("새벽 2시 “엄마… 인누와~” / 엄마 눈 밑이 판다 / 아빠가 대신 갔다가 “아니! 엄마!! 인누와!!” → 아빠 탈락", "‘인누와’도 곧 사라질 말 / “엄마 피곤하지? 이번엔 내가 토닥토닥 해 줄게.” — 그날 밤은 엄마가 먼저 잠들었다"),
    20: ("“씻을 시간이야.” “거짓말! 아니야!” “시계가 거짓말이야!” (그림을 더 그리고 싶어서) / 물이 눈에 들어갈까 봐 부르르 떠는 태오도 “거짓말! 아니야! 안 무서워!!” / 마지막엔 “김태오!! 그건 내 말이야!!”", "‘아니야’ 뒤에 숨은 마음을 말하는 재이 / 잠수를 배운 누나가 “태오도 할 수 있어. 누나가 손 잡아 줄게.” / 수건 + 레몬 냄새 + 누나 손으로 “조금은 안 무서웠어!”"),
    21: ("아침 7시 30분 “이 안 닦아! 싫어!!” 도망, 식탁·소파 밑을 기어서 숨기 / 저녁에도 또 도망 → 엄마 아빠 판다 눈 / “나와라 타노스! 어벤져스 칫솔!” → 아빠 “거품 공격에 졌다!” / 다음 날 아침 또 도망", "“누나는 누나니까!”라던 태오가 타노스 놀이로 이를 다 닦는다 / 그래도 엄마 아빠는 아침에도, 저녁에도 오늘도 타노스가 된다"),
    22: ("“어벤져스 캐릭터로 싸우자! 치앗촤!” 감독 태오가 역할 정하기 / 누나는 블랙 위도우, 엄마는 핀타스틱포 쑤우, 둘 다 “예뻐!” / 아빠 타노스 “아임 인에비터블” → 오늘도 패배", "예뻐서 지키겠다던 태오에게 누나와 엄마가 “우리도 싸울 수 있어!” — 서로 지켜 주는 가족"),
    7: ("“오늘 티비 보는 날이야?” “어린이집 안 가는 날이야?” “맛있는 거 죠!” 매일 똑같은 질문 / 아빠가 만든 표를 태오는 못 읽는다", "“엄마 아빠랑 노는 게 제일 맛있어.”"),
    6: ("블루베리·스테잌로저·미스테리 / 우엉조림을 보고 “우엉은 마법사잖아!” / 할머니께 “쓰구미!” / 아기 재이의 “짐방!” (공책엔 꿍끙이 = 뿡뿡이도)", "“지금만 들을 수 있는 보물 같은 말” — 누나가 만들어 주는 태오어 사전"),
}

hdr_fill = PatternFill("solid", fgColor="F0609D")
sub_fill = PatternFill("solid", fgColor="FFF4E4")
thin = Side(style="thin", color="E6CDB8")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="top", wrap_text=True)

def sheet(ws, headers, widths, rows, title=None):
    r0 = 1
    if title:
        ws.cell(1, 1, title).font = Font(bold=True, size=14, color="3B2A24")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        r0 = 3
    for i, h in enumerate(headers, 1):
        c = ws.cell(r0, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hdr_fill
        c.alignment = center
        c.border = border
        ws.column_dimensions[c.column_letter].width = widths[i - 1]
    for r, row in enumerate(rows, r0 + 1):
        for i, v in enumerate(row, 1):
            c = ws.cell(r, i, v)
            c.alignment = center if isinstance(v, (int, float)) else wrap
            c.border = border
            if r % 2 == 0:
                c.fill = sub_fill
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    ws.auto_filter.ref = f"A{r0}:{ws.cell(r0 + len(rows), len(headers)).coordinate}"
    return r0

wb = Workbook()
ws = wb.active
ws.title = "등장인물"
rows = []
for i, c in enumerate(d["CHARACTERS"], 1):
    look, clothes = LOOK[c["id"]]
    rows.append([i, c["name"], c["tag"], look, clothes, c["line"]])
sheet(ws, ["순서", "이름", "한 줄 소개", "생김새 (그림 기준)", "옷", "성격·특징"], [6, 10, 26, 34, 34, 60], rows,
      "「오늘도 우리 집」 등장인물 — 재이 · 할머니 · 엄마 · 아빠 · 할아버지 · 태오")

ws = wb.create_sheet("에피소드")
rows = []
for ep in d["EPISODES"]:
    ha, heart = POINTS[ep["id"]]
    cast = []
    for p in ep["panels"]:
        for ch in p.get("chars", []):
            n = NAMES[ch["c"]]
            if n not in cast:
                cast.append(n)
    order = [c["name"] for c in d["CHARACTERS"]]
    cast.sort(key=order.index)
    rows.append([ep["id"], ep["title"], ep["summary"], ", ".join(cast), ha, heart, ep["lesson"], ep["talk"], len(ep["panels"]), "완성"])
first = sheet(ws, ["화", "제목", "줄거리", "등장인물", "웃음 포인트", "감동 포인트", "오늘의 한 줄 (교훈)", "가족과 이야기해 봐요", "컷 수", "상태"],
              [5, 18, 44, 22, 36, 36, 32, 34, 7, 9], rows, "에피소드 목록")
n = len(rows)
for k in range(3):  # 다음 화를 적을 빈 줄
    r = first + n + 1 + k
    ws.cell(r, 1, n + 1 + k).alignment = center
    ws.cell(r, 10, "기획 중")
    for i in range(1, 11):
        ws.cell(r, i).border = border
dv = DataValidation(type="list", formula1='"아이디어,기획 중,콘티,완성"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"J{first + 1}:J{first + n + 20}")

ws = wb.create_sheet("컷 대본")
rows = []
for ep in d["EPISODES"]:
    for i, p in enumerate(ep["panels"], 1):
        kind = p.get("bg", "plain").split(":")[0]
        place = BG.get(kind, kind)
        if any(fp[0] == "phoneframe" for fp in p.get("fp", [])):
            place = "영상통화 화면"
        cast = []
        for ch in p.get("chars", []):
            n = NAMES[ch["c"]]
            if n not in cast:
                cast.append(n)
        narr = " / ".join(x for x in [p.get("n"), p.get("n2")] if x)
        lines = "\n".join(("(" + b["k"] + ") " if b.get("k") in ("think", "whisper") else "") + b["t"] for b in p.get("b", []))
        lines = lines.replace("(think) ", "(생각) ").replace("(whisper) ", "(속삭임) ")
        sfx = ", ".join(f["t"] for f in p.get("sfx", []))
        rows.append([ep["id"], ep["title"], i, place, ", ".join(cast), narr, lines, sfx])
sheet(ws, ["화", "제목", "컷", "장소·효과", "등장인물", "내레이션", "대사", "효과음"], [5, 18, 5, 14, 22, 36, 52, 16], rows,
      "컷별 대본 (웹툰 화면과 같은 순서)")

ws = wb.create_sheet("태오어 사전")
rows = [[i, w["say"], w["real"], w["what"], "", ""] for i, w in enumerate(d["WORDS"], 1)]
rows += [[len(rows) + i, "", "", "", "", ""] for i in range(1, 21)]
sheet(ws, ["번호", "태오가 하는 말", "원래 말", "무엇인지", "처음 들은 날", "메모 (상황·에피소드 아이디어)"], [6, 18, 18, 22, 14, 44], rows,
      "태오어 사전 — 새로 들은 말을 빈 줄에 적어 주세요")

ws = wb.create_sheet("우리 집 말 모음")
rows = [[i, "태오 입버릇", w["say"], "", w["when"], ""] for i, w in enumerate(d["SAYINGS"], 1)]
rows += [[len(rows) + i, "재이 입버릇", w["say"], "", w["when"], ""] for i, w in enumerate(d["JSAYINGS"], 1)]
rows += [[len(rows) + i, "재이 아기 말", w["say"], w["real"], w["what"], ""] for i, w in enumerate(d["JWORDS"], 1)]
rows += [[len(rows) + i, "", "", "", "", ""] for i in range(1, 16)]
r0 = sheet(ws, ["번호", "누구", "한 말", "원래 말", "언제·사연", "메모"], [6, 14, 26, 14, 50, 30], rows,
      "우리 집 말 모음 — 태오 입버릇, 재이가 아기 때 한 말 (빈 줄에 더 적어 주세요)")
dv3 = DataValidation(type="list", formula1='"태오 입버릇,재이 입버릇,재이 아기 말,태오어,기타"', allow_blank=True)
ws.add_data_validation(dv3)
dv3.add(f"B{r0 + 1}:B{r0 + 60}")

ws = wb.create_sheet("다음 화 아이디어")
ideas = [
    ["다리가 부러진 장난감", "태오", "(가족 제보) 태오는 다리가 부러진 장난감은 잘 안 가지고 논다 → 11화로 만듦", "", "완성"],
    ["태권도 승급 심사", "태오", "긴장한 태오가 심사장에서 넘어진다. 가족 응원석에선 할아버지가 제일 크게 “메가랩터킥!”을 외치고…", "넘어져도 다시 일어나면 그게 진짜 실력이에요.", "아이디어"],
    ["할머니의 파마 대작전", "할머니, 재이", "재이가 미용실 놀이로 할머니 머리를 ‘펴’ 드리겠다고 나섰다. 결과는 더 뽀글뽀글?", "있는 그대로의 모습이 가장 예뻐요.", "아이디어"],
    ["아빠의 비밀 요리", "아빠, 태오", "엄마 생일, 아빠와 태오가 몰래 미역국을 끓인다. 미역이 냄비 밖으로 넘쳐 나오는데…", "마음을 담으면 서툴러도 최고의 선물이에요.", "아이디어"],
    ["누나 대신 태오", "태오, 재이", "누나가 아파서 학교에 못 간 날, 태오가 누나의 ‘하루 비서’가 된다.", "도움을 받던 사람도 도움을 줄 수 있어요.", "아이디어"],
    ["", "", "", "", ""],
    ["", "", "", "", ""],
]
r0 = sheet(ws, ["제목 (안)", "주인공", "줄거리 한 줄", "오늘의 한 줄", "상태"], [24, 16, 62, 36, 10], ideas,
           "다음 화 아이디어 — 가족이 같이 채워 보세요")
dv2 = DataValidation(type="list", formula1='"아이디어,기획 중,콘티,완성"', allow_blank=True)
ws.add_data_validation(dv2)
dv2.add(f"E{r0 + 1}:E{r0 + 40}")
for w in wb.worksheets:
    w.sheet_view.zoomScale = 100
wb.save(out)
print("saved", out)
